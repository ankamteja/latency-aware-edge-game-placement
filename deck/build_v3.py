"""Build v3 of the team deck from v2: held-out replicate, two live-controller
slides, limits and conclusion brought up to date.

usage: python3 -I build_v3.py <v2.pptx> <out.pptx> <figures_dir>
"""
import copy
import sys
from pathlib import Path

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Emu

src, dst, figs = sys.argv[1], sys.argv[2], Path(sys.argv[3])
prs = Presentation(src)
S = prs.slides


def shape(slide, name):
    return next(s for s in slide.shapes if s.name == name)


def set_text(slide, name, text):
    """Replace a single-paragraph box's text, keeping the first run's format."""
    tf = shape(slide, name).text_frame
    p = tf.paragraphs[0]
    for extra in tf.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)
    runs = p.runs
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def set_paras(slide, name, heading, bullets):
    """Heading paragraph + bullet paragraphs, cloned from the box's own
    first and second paragraphs so fonts, colours and spacing are kept."""
    tf = shape(slide, name).text_frame
    paras = tf.paragraphs
    head_xml = copy.deepcopy(paras[0]._p)
    bullet_xml = copy.deepcopy(paras[1]._p if len(paras) > 1 else paras[0]._p)
    body = paras[0]._p.getparent()
    for p in paras:
        body.remove(p._p)
    for xml, text in [(head_xml, heading)] + [(bullet_xml, "▸ " + b) for b in bullets]:
        node = copy.deepcopy(xml)
        body.append(node)
    for p, text in zip(tf.paragraphs, [heading] + ["▸ " + b for b in bullets]):
        p.runs[0].text = text
        for r in p.runs[1:]:
            r._r.getparent().remove(r._r)


def remove(slide, *names):
    for n in names:
        el = shape(slide, n)._element
        el.getparent().remove(el)


def resize(slide, name, left=None, top=None, width=None, height=None):
    s = shape(slide, name)
    if left is not None: s.left = Emu(left)
    if top is not None: s.top = Emu(top)
    if width is not None: s.width = Emu(width)
    if height is not None: s.height = Emu(height)


def picture(slide, png, left, top, width=None, height=None):
    slide.shapes.add_picture(str(figs / png), Emu(left), Emu(top),
                             Emu(width) if width else None, Emu(height) if height else None)


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def grid(slide, first_box, cols, rows_text):
    """The tables are grids of rectangles + text boxes named 'TextBox k'.
    first_box is the header's first text box number; boxes step by 2."""
    k = first_box
    for row in rows_text:
        for cell in row:
            set_text(slide, f"TextBox {k}", cell)
            k += 2



def clone_after(idx):
    """A copy of slide idx (shapes except pictures), placed right after it."""
    s = prs.slides[idx]
    new = prs.slides.add_slide(s.slide_layout)
    for sh in list(new.shapes):
        sh._element.getparent().remove(sh._element)
    bg = s._element.cSld.find(qn("p:bg"))
    if bg is not None:                  # the dark background is set per slide
        new._element.cSld.insert(0, copy.deepcopy(bg))
    for sh in s.shapes:
        if sh.shape_type != 13:
            new.shapes._spTree.append(copy.deepcopy(sh._element))
    lst = prs.slides._sldIdLst
    el = lst[-1]
    lst.remove(el)
    lst.insert(idx + 1, el)
    return new


# ---- slide 1: authors (full registered names as they arrive) ----------------
AUTHORS = ["Ankam Charan Teja", "Aranya Roy", "Nikhil Kumar Reddy", "Yashvardhan Deep", "Akshit Singh"]
set_text(S[0], "TextBox 6", "  ·  ".join(AUTHORS))

# ---- slide 5: topology diagram redrawn with the /30 subnets the YAML configures
pic = next(sh for sh in S[4].shapes if sh.shape_type == 13)
pos = (pic.left, pic.top, pic.width, pic.height)
pic._element.getparent().remove(pic._element)
S[4].shapes.add_picture(str(Path(__file__).resolve().parent / "assets" / "topology_30.png"), *pos)

# ---- slide 9: held-out replicate ---------------------------------------------
s = S[8]
set_text(s, "TextBox 3", "Three latency models, each scored on replicates it was not fitted on")
set_paras(s, "TextBox 5", "Dataset and models", [
    "Dataset: 150 cells, 450 rows of 30 s windows (4 grid replicates + 2 deployment rounds)",
    "Analytical: RTT + size/bandwidth, no load term",
    "Queueing (M/M/1): RTT + s/(1−u), s fitted = 2.4 ms",
    "Learned: RTT + gradient-boosted trees on utilization, players, cores",
    "Fit on one replicate, test on the other: median abs error 5.8 / 1.8 / 0.9 ms (analytical / queueing / learned)",
    "Fit on both, test on two held-out replicates: 6.9 / 2.2 / 1.1 ms; F1 at 50 ms 0.00 / 0.60 / 0.86",
    "Held out, learned vs queueing is within noise (95% interval -4.2 to +0.1 ms)",
])
notes(s, "The dataset is the artifact: every row is a measured window, not a "
         "formula output. The first error line fits on one replicate and tests "
         "on the other. The second fits on both and tests on two more replicates "
         "measured hours later and never used in a fit. The ranking holds there, "
         "but learned vs queueing is not a significant difference on those 50 cells. F1 is for the decision 'will this node break 50 ms at "
         "this load': the analytical model never predicts a violation, so its "
         "accuracy looks fine and its F1 is zero. Source: 08-results 8.4.")

# ---- slide 10: placement deployed (ratio and wording checked) -----------------
set_paras(S[9], "TextBox 5", "Worst player's measured p95, round 1 / round 2", [
    "40 players: all three 26 / 16 ms",
    "80: analytical 29.9 / 28.2 s · queueing 20 / 20 ms · learned 20 / 20 ms",
    "120: analytical 17.8 / 18.7 s · queueing 31 / 37 ms · learned 32 / 32 ms",
    "160: analytical 20.0 / 28.3 s · queueing 35 / 40 ms · learned 39 / 55 ms",
    "Graded by its own formula, as a simulator would, the load-blind placement looks best (5–15 ms)",
    "Deployed: 500–1500x worse; the overloaded edge1 held only 43–71 of its 80–100 players",
    "Learned predicts best, but places almost the same as fitted M/M/1",
])

# ---- new slides 11 and 12: the live controller -------------------------------
live = clone_after(9)
set_text(live, "TextBox 2", "Live Controller on the Real Testbed")
set_text(live, "TextBox 3", "Same 13-minute schedule of players under each controller; one match moved at most every 10 s")
set_paras(live, "TextBox 5", "Worst-off players, median of 3 runs each", [
    "Load-blind formula: 1.1 s (76 ms to 4.5 s), over 100 ms 53% of the time, never moves a match",
    "Queueing (M/M/1): 99 ms (72 to 166), 50%: loads edge2 to 80 players, past its knee, in every run",
    "Learned model: 39 ms (33 to 47), 33%: stops edge2 at 70 and uses edge3, in every run",
    "Learned, corrected live: 33 ms (32 to 34), 8%, with 15 moves: more caution, fewer bad moments",
    "10–16 moves for the load-aware controllers (load-blind: 0), 6 s to reconnect",
    "Learned's remaining time over budget: edge1 at 40 players, just below its knee",
])
picture(live, "fig6_live_ab.png", 5650000, 1700000, width=6240000)
notes(live, "This is the placement result in closed loop. A controller re-runs "
            "placement every 10 s on live measurements and moves whole matches. "
            "Each controller ran three times; the chart shows each one's median run. "
            "The load-blind range is wide because in one run edge1's players most "
            "likely timed out, which lowered the latency of those left. "
            "Both runs see exactly the same arrivals and departures; new matches "
            "land on the nearest node, as a distance-based matchmaker would do. "
            "The formula sees 5 ms on edge1 at any load, so it never acts. "
            "Source: docs/09-live.md 9.6, results/live/.")

drop = clone_after(10)
set_text(drop, "TextBox 2", "A Node Gets Slower and No Model Is Told")
set_text(drop, "TextBox 3", "edge1's CPU cut from 0.5 to 0.25 cores at 220 s; the controller only sees measurements")
set_paras(drop, "TextBox 5", "After the cut", [
    "Queueing and offline learned keep 40 players on the slowed edge1 (6 runs)",
    "They are over 100 ms 99–100% of the time after the cut",
    "Corrected live, with memory: 6% after the cut (4 to 13%, 3 runs), 8–12 moves",
    "Caveat: it learns from players logging in, so in 2 of 3 runs edge1 was already near empty at the cut",
    "Clean detection: 30 players at the cut, a match moved off 27 s later",
])
picture(drop, "fig7_live_capacity_drop.png", 5650000, 1550000, height=5000000)
notes(drop, "The node got slower and nothing told the controller: a noisy neighbour, a throttled host, a smaller VM. The offline model keeps using the curve measured at 0.5 cores. The live-corrected model compares every 10 s window with its prediction, estimates how much more each player now costs edge1, and remembers loads it measured over budget for 5 minutes. Be honest if asked: the memory was added after a first run in which the model noticed, then forgot and moved a match back. And the correction also learns from the seconds in which a new match is logging in, so in two of its three runs it already distrusted edge1 and had moved almost everyone off before the cut. Only one run, with 30 players on edge1 at the cut, shows it detecting the change: it moved a match off 27 s later. Source: docs/09-live.md 9.6.")

# ---- limitations (now slide 13) ------------------------------------------------
s = S[12]
set_paras(s, "TextBox 4", "Scope", [
    "One host: bootstrap intervals cover cells, not whole replicates (only four)",
    "Capacities and delays were chosen to create the trade-off",
    "Bots, not humans; hostile mobs disabled",
    "A move is a reconnect (about 6 s), not a state handover",
    "Live: 3 runs per controller, wide ranges; direction, not effect sizes",
])
set_paras(s, "TextBox 6", "Measurement", [
    "Knees are unstable: replicates differ 6–80x there",
    "Host power mode changes capacity 2–3x: all clean runs in Performance mode",
    "edge3 and cloud never pushed past their knee (100-player ceiling)",
    "Bandwidth term negligible: links carry no background traffic",
])
notes(s, "Every number on the previous slides comes with these limits. The "
         "power-mode effect is the most interesting one: with nothing else "
         "changed, the laptop's Balanced mode made every node 2-3x slower and "
         "moved edge1's knee from 50 players to under 20. No latency formula "
         "sees that; a measured model does. Source: 08-results 8.6.")

# ---- conclusion (now slide 14) -------------------------------------------------
s = S[13]
set_paras(s, "TextBox 4", "What this work shows", [
    "The load-blind formula is right only below a node's saturation point",
    "Above it the error is seconds: 18–30 s worst-player latency, deployed",
    "Load-aware placement keeps the same players at 16–55 ms",
    "Live, the formula never moves a match; a load-aware controller keeps players near 33 ms",
    "A learned model predicts best; a fitted M/M/1 already places right here",
])
set_paras(s, "TextBox 6", "Future work", [
    "More replicates on a dedicated host, intervals on the knees",
    "State handover instead of reconnect; place new matches with the model",
    "Jitter, loss and bandwidth caps; mob AI back as a second load source",
])
notes(s, "Close on the deployed and live numbers, they are the strongest "
         "evidence. Repo, raw data, dataset and the live dashboard: "
         "github.com/ankamteja/latency-aware-edge-game-placement.")

# ---- final pass over every slide: plain wording, no em dashes, no stale numbers
FIX = [
    ("10-minute slot — one idea per slide.", "10-minute slot, one idea per slide."),
    ("Good [morning/afternoon]. ", ""),
    ("A navigation analogy — routing blind to traffic", "A navigation analogy: routing blind to traffic"),
    ("🗺  ", ""), ("🖧  ", ""),
    ("The shortest road is not always the fastest — traffic matters.",
     "The shortest road is not always the fastest; traffic matters."),
    ("The Load-Blind Formula — and Why Simulators Hide the Error",
     "The Load-Blind Formula, and Why Simulators Hide the Error"),
    ("Both terms are network-only — server is assumed instantaneous",
     "Both terms are network-only: the server is assumed to answer instantly"),
    ("A loaded server queues every game tick — that delay is absent",
     "A busy server makes every action wait for its game tick; the formula has no term for that"),
    ("queues work — every player action", "queues work: every player action"),
    ("under load — and can a predictor", "under load, and can a predictor"),
    ("Client + three edge nodes + cloud — all on one physical host",
     "A client, three edge nodes and a cloud node, all on one laptop"),
    ("All four probes run in parallel — same load, same moment",
     "All four probes run in parallel: same load, same moment"),
    ("mineflayer library — actual game protocol, not HTTP.",
     "mineflayer library, which speaks the real game protocol, not HTTP."),
    ("Decomposing where delay lives — network, OS, process, or game thread",
     "Where the delay comes from: network, OS, server process or game thread"),
    ("Kernel to kernel — network only", "Network only"),
    ("ICMP flat ⟹ the network is not the cause.", "If ICMP stays flat, the network is not the cause."),
    ("edge2 (1 core) is flat to 60 players, breaks between 80 and 100",
     "edge2 (1 core) is flat to 60 players; knee at 80, breaks by 100"),
    ("Containerlab emulates links with exact RTT budgets",
     "netem adds each link's delay; the RTT is measured with ping"),
    ("Both runs see exactly the same arrivals and departures",
     "Every run sees exactly the same arrivals and departures"),
    ("edge1 at 50 players: MSPT 6.5 ms", "edge1 at 50 players: MSPT 6.6 ms"),
    ("Every ms above probe 1 is model error — unmeasured queueing.",
     "Every ms above probe 1 is delay the formula cannot see."),
]
for sl in prs.slides:
    frames = [sh.text_frame for sh in sl.shapes if sh.has_text_frame]
    if sl.has_notes_slide:
        frames.append(sl.notes_slide.notes_text_frame)
    for tf in frames:
        for para in tf.paragraphs:
            for run in para.runs:
                for a, b in FIX:
                    if a in run.text:
                        run.text = run.text.replace(a, b)
# slide 7: cells for loads edge1 was never run at
for name in ("TextBox 57", "TextBox 59", "TextBox 67", "TextBox 69"):
    set_text(S[6], name, "not run")
notes(S[5], "All four probes run at the same time, on the same link, under the same "
            "load, so the delay can be traced to a layer instead of guessed. Probe 1 "
            "is what the formula predicts; anything above it is delay the formula "
            "cannot see. p95 means 95% of messages were faster than this value: it "
            "shows what the unluckiest players feel, not the average. Measured on "
            "edge1 from 1 to 60 players: ICMP stayed at 5.1 to 5.3 ms, the status "
            "ping moved only at 60 players (16 to 32 ms), and the in-game p95 went "
            "from 6 ms to 6.7 to 8.6 seconds. Source: docs/08-results.md 8.2.")
notes(S[0], "This talk is about a mismatch between how edge placement algorithms "
            "choose a server node and what players actually experience when that node "
            "gets busy. Repository and a replay of the live dashboard: "
            "github.com/ankamteja/latency-aware-edge-game-placement. 10 minutes, one "
            "idea per slide.")
set_paras(S[13], "TextBox 4", "What this work shows", [
    "The load-blind formula is right only below a node's saturation point",
    "Above it the error is seconds: 18–30 s worst-player latency, deployed",
    "Load-aware placement keeps the same players at 16–55 ms",
    "Live, 3 runs each: the formula never moves a match; the learned controller's median is 39 ms",
    "Utilization alone is not enough: live, M/M/1 overloaded edge2 in all 3 runs, the learned model in none",
])

for i, sl in enumerate(prs.slides, 1):
    texts = [sh.text_frame.text for sh in sl.shapes if sh.has_text_frame]
    if sl.has_notes_slide:
        texts.append(sl.notes_slide.notes_text_frame.text)
    for t in texts:
        assert "—" not in t and "[" not in t.replace("[1]", ""), (i, t[:80])

prs.save(dst)
print("wrote", dst)
