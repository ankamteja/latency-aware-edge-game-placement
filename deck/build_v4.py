"""v4: the v3 deck arranged in the report's sections (Abstract, Introduction,
Methodology, Implementation, Results, Conclusion and Future Work): an
Abstract and an Implementation slide are added and every slide carries its
section label.

usage: python3 -I build_v4.py <v3.pptx> <out.pptx>
"""
import copy
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

src, dst = sys.argv[1], sys.argv[2]
prs = Presentation(src)
TEAL = RGBColor(0x00, 0xC2, 0xB2)


def shape(slide, name):
    return next(s for s in slide.shapes if s.name == name)


def set_text(slide, name, text):
    tf = shape(slide, name).text_frame
    p = tf.paragraphs[0]
    for extra in tf.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)
    p.runs[0].text = text
    for r in p.runs[1:]:
        r._r.getparent().remove(r._r)


def set_paras(slide, name, heading, bullets):
    tf = shape(slide, name).text_frame
    paras = tf.paragraphs
    head_xml = copy.deepcopy(paras[0]._p)
    bullet_xml = copy.deepcopy(paras[1]._p if len(paras) > 1 else paras[0]._p)
    body = paras[0]._p.getparent()
    for p in paras:
        body.remove(p._p)
    for xml in [head_xml] + [bullet_xml] * len(bullets):
        body.append(copy.deepcopy(xml))
    for p, text in zip(tf.paragraphs, [heading] + ["▸ " + b for b in bullets]):
        p.runs[0].text = text
        for r in p.runs[1:]:
            r._r.getparent().remove(r._r)


def clone_to(idx, pos):
    """A copy of slide idx (pictures left out), moved to position pos."""
    s = prs.slides[idx]
    new = prs.slides.add_slide(s.slide_layout)
    for sh in list(new.shapes):
        sh._element.getparent().remove(sh._element)
    bg = s._element.cSld.find(qn("p:bg"))
    if bg is not None:
        new._element.cSld.insert(0, copy.deepcopy(bg))
    for sh in s.shapes:
        if sh.shape_type != 13:
            new.shapes._spTree.append(copy.deepcopy(sh._element))
    lst = prs.slides._sldIdLst
    el = lst[-1]
    lst.remove(el)
    lst.insert(pos, el)
    return new


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


S = list(prs.slides)
research_q, limitations = S[3], S[12]

# ---- Implementation: after the measurement ladder (v3 slide 6) --------------
impl = clone_to(12, 6)
set_text(impl, "TextBox 2", "Implementation")
set_paras(impl, "TextBox 4", "Testbed and measurement", [
    "containerlab topology, tc netem delays, cgroup cpu.max quotas on disjoint cores",
    "Paper 1.21.4 servers; topology/verify.sh checks links, RTT, CPU limits, power mode",
    "bots/bot.js: mineflayer bots, 10 per process; probe.js: TCP and status ping",
    "collect.sh, grid.sh: one cell = settle, baseline, joins, 90 s steady window",
])
set_paras(impl, "TextBox 6", "Analysis and live system", [
    "build_dataset.py: 150 cells, 450 rows of 30 s windows",
    "models.py: three models behind one predict(node, players)",
    "experiment.py: exhaustive placement search; deploy_plan.sh deploys it",
    "live/: 2 s controller loop, browser dashboard, 29 unit tests",
    "rebuild.sh: every table and figure rebuilt from the raw data",
])
notes(impl, "How it is built. Everything is scripts in the repository: the "
            "topology and its checks, the bots and probes, the cell collector, "
            "the dataset builder, the three models behind one interface so every "
            "algorithm runs with any of them, the placement search and deployer, "
            "and the live controller with its dashboard. One command rebuilds "
            "every number in the report from the raw measurements. Source: "
            "README, report Section III.")

# ---- Abstract: right after the title ------------------------------------------
abst = clone_to(3, 1)
set_text(abst, "TextBox 2", "Abstract")
set_text(abst, "TextBox 4",
         "Edge placement algorithms choose a node with a latency formula that has "
         "no term for server load, and are usually evaluated in simulators built on "
         "the same formula. We measure the error on a real testbed: four CPU-capped "
         "game servers, protocol-level bots, latency measured at four depths at once, "
         "and each model's placement deployed for real and run as a live controller.")
set_paras(abst, "TextBox 6", "Key results", [
    "Network delay stays within 0.3 ms at any load; the smallest node's p95 goes from 15 ms to 6.7–8.6 s",
    "Deployed, the load-blind placement leaves the worst players at 18–30 s; load-aware placements 16–55 ms",
    "The learned model predicts best (0.9 ms median error, 1.1 ms held out) but places like a fitted M/M/1",
    "Live, 3 runs each: formula 1.1 s, M/M/1 99 ms, learned 39 ms, learned corrected live 33 ms",
])
notes(abst, "The whole talk in one slide: the formula placement algorithms use "
            "has no load term, so on a busy edge node it is wrong by seconds, "
            "and a simulator built on it cannot show that. Every number here is "
            "measured, and each comes back with its detail on a later slide.")

# ---- section label on every slide but the title -----------------------------
SECTIONS = (["ABSTRACT"] + ["I · INTRODUCTION"] * 3 + ["II · METHODOLOGY"] * 2
            + ["III · IMPLEMENTATION"] + ["IV · RESULTS"] * 6
            + ["V · CONCLUSION AND FUTURE WORK"] * 2)
slides = list(prs.slides)
assert len(slides) == 1 + len(SECTIONS), len(slides)
W = prs.slide_width
for sl, label in zip(slides[1:], SECTIONS):
    box = sl.shapes.add_textbox(Emu(W - 274320 - 4572000), Emu(60000), Emu(4572000), Emu(300000))
    p = box.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    r = p.add_run()
    r.text = label
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = TEAL

for i, sl in enumerate(prs.slides, 1):
    texts = [sh.text_frame.text for sh in sl.shapes if sh.has_text_frame]
    if sl.has_notes_slide:
        texts.append(sl.notes_slide.notes_text_frame.text)
    for t in texts:
        assert "—" not in t, (i, t[:80])

cp = prs.core_properties
cp.title = "Latency-Aware Edge Game Server Placement"
cp.author = "Ankam Charan Teja, Aranya Roy, Nikhil Kumar Reddy, Yashvardhan Deep, Akshit Singh"
cp.last_modified_by = ""
prs.save(dst)
print("wrote", dst, len(prs.slides), "slides")
